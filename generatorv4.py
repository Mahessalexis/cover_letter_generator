"""
Générateur de Lettres de Motivation - Révélation Entreprendre
Version 4.0 - Multilingue (FR/EN) et Multi-domaines
"""

import os
import io
import re
from datetime import datetime
from flask import Flask, request, render_template_string, session, send_file, jsonify
from dotenv import load_dotenv
from openai import OpenAI
from docx import Document

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "change-moi-en-prod-avec-une-vraie-cle-secrete")

# ============================================================================
# CONFIGURATION
# ============================================================================
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
RE_SECRET_CODE = os.getenv("RE_SECRET_CODE", "REVELATION2025")

# Template Word (même fichier pour FR et EN, seul le contenu change)
TEMPLATE_DOCX_PATH = "template placeholder.docx"

# ============================================================================
# LETTRES MODÈLES (Style de référence)
# ============================================================================

TEMPLATE_LETTER_FR = """
Actuellement en {academic_status} et fort de plusieurs expériences professionnelles, je suis à la recherche d'un {internship_type} au sein de {company_name}. Ce qui m'attire chez {company_name}, c'est son positionnement de référence et sa réputation d'excellence dans son domaine. Mon échange avec des professionnels du secteur m'a confirmé que le dynamisme et la qualité des équipes favorisent une réelle montée en compétences, ce que je recherche particulièrement pour mon stage.

Ma volonté de travailler dans ce domaine s'est affirmée lors de mes précédentes expériences professionnelles. Les retours sur la courbe d'apprentissage et la technicité du métier m'ont vivement intéressé. Curieux d'en découvrir davantage, j'ai multiplié les initiatives pour approfondir mes connaissances. Cette expérience m'a donné un premier aperçu de la rigueur technique requise et a confirmé mon désir de poursuivre dans cette voie.

J'ai toujours cultivé une recherche de l'excellence, tant sur le plan professionnel qu'académique. Classé dans le top 5% de ma promotion à l'EDHEC, je trouve également le temps de m'investir dans l'association étudiante Révélation Entreprendre pour organiser le plus grand concours de start-ups étudiant, réunissant plus de 500 startups candidates et doté de 350 000€. Ce poste m'a permis de développer des compétences en leadership, en gestion de projet et en travail d'équipe.

Je me tiens à votre disposition pour un entretien afin d'échanger davantage sur mon parcours, mes motivations et la manière dont je pourrai apporter ma contribution au travail de vos équipes.

Dans l'attente du plaisir de vous rencontrer, je vous prie d'agréer, Madame, Monsieur, l'expression de mes respectueuses salutations.
""".strip()

TEMPLATE_LETTER_EN = """
I am currently a Master in Finance student at EDHEC Business School. After gaining experience in M&A, transaction services and audit, I am looking for an internship in M&A at DC Advisory in London.

What attracts me to DC Advisory is its strong reputation in the mid-market segment and its ability to deliver high-quality, tailor-made advice on complex cross-border transactions. The firm's emphasis on collaboration across offices and its exposure to international clients make it an ideal environment to grow as an M&A analyst. The combination of technical excellence and human values perfectly matches what I am looking for. These are also values I have developed through playing chess — a discipline in which I competed in the French Championships and learned focus, discipline, strategic thinking, and even team spirit when playing for one of the top youth teams in France. I am also drawn to DC Advisory for the trust it inspires among the companies it advises. One recent example that particularly caught my attention is DC Advisory's role as exclusive financial advisor to Antin Infrastructure Partners on its acquisition of AquaVista Watersides & Marinas (announced Sept 2025). This transaction demonstrates the firm's trusted advisory role in infrastructure and its ability to deliver value in highly specialised, cross-border infrastructure deals.

My motivation to work in M&A grew stronger during my last internship in audit at Advolis Orfis, where I interacted with colleagues from the Transaction Services team. Their feedback on the steep learning curve and the technical nature of the work greatly interested me. Eager to learn more, I asked to take part in Transaction Services projects alongside my daily work. For two months, I combined audit tasks with due diligence assignments in consumer and retail, often working evenings to be fully involved. This experience gave me a first taste of the technical rigor of mergers and acquisitions and confirmed my ambition to follow transactions throughout their entire cycle — from negotiation to closing. Building on this first experience, I will further strengthen my technical skills during my exchange semester at Bocconi University, while simultaneously interning in M&A at the small-cap advisory firm TransferiuS from January to June 2026.

I have always pursued excellence, both professionally and academically. Ranked in the top 5% of my class at EDHEC, I also find the time to lead a 10-person team in the student association « Révélation Entreprendre » to organize the largest student start-up competition, bringing together more than 500 startups and offering €350,000 in prizes. This role allowed me to develop leadership, project management, and teamwork skills, which I believe would be valuable in a fast-paced environment like DC Advisory.

I remain at your disposal for an interview to further discuss my background, motivations, and how I could contribute to your team's work.

Yours sincerely,
""".strip()


# ============================================================================
# CONSTRUCTION DU PROMPT OPTIMISÉ
# ============================================================================
def build_prompt(form):
    """
    Construit un prompt optimisé avec support multilingue et multi-domaines.
    """
    # Récupération des données
    language = form.get("language", "fr").strip()
    company_name = form.get("company_name", "").strip()
    internship_type = form.get("internship_type", "").strip()
    availability = form.get("availability", "").strip()
    previous_experience = form.get("previous_experience", "").strip()
    networking_info = form.get("networking_info", "").strip()
    top_qualities = form.get("top_qualities", "").strip()
    striking_fact = form.get("striking_fact", "").strip()
    future_plan = form.get("future_plan", "").strip()
    re_role = form.get("re_role", "").strip()
    academic_status = form.get("academic_status", "M1 au Finance Track de l'EDHEC" if language == "fr" else "Master in Finance at EDHEC Business School").strip()
    industry_sector = form.get("industry_sector", "").strip()
    location = form.get("location", "").strip()

    # Sélection du template selon la langue
    template_letter = TEMPLATE_LETTER_FR if language == "fr" else TEMPLATE_LETTER_EN

    # Instructions conditionnelles selon la langue
    if language == "fr":
        # FRANÇAIS
        networking_instruction = ""
        if networking_info:
            networking_instruction = f"""
   - INTÈGRE OBLIGATOIREMENT cette information de networking : "{networking_info}"
   - Utilise-la pour justifier ton intérêt par des retours internes concrets"""
        else:
            networking_instruction = """
   - Pas de contact mentionné : fais référence à la réputation de l'entreprise ou à des sources publiques"""

        fact_instruction = ""
        if striking_fact:
            fact_instruction = f"""
   - MENTIONNE OBLIGATOIREMENT ce fait marquant : "{striking_fact}"
   - Intègre-le naturellement pour montrer ta connaissance de l'entreprise"""

        future_instruction = ""
        if future_plan:
            future_instruction = f"""
   - MENTIONNE ton projet futur : "{future_plan}" pour montrer ta curiosité et ton ambition"""

        sector_instruction = ""
        if industry_sector:
            sector_instruction = f"""
   - Adapte le vocabulaire et les références au secteur : {industry_sector}"""

        prompt = f"""Tu es un rédacteur expert en lettres de motivation professionnelles.

OBJECTIF : Rédiger le CORPS d'une lettre de motivation EN FRANÇAIS (de "Madame, Monsieur," jusqu'à la formule de politesse incluse).

═══════════════════════════════════════════════════════════════════════════════
MODÈLE DE RÉFÉRENCE (imite ce style précisément)
═══════════════════════════════════════════════════════════════════════════════
{template_letter}
═══════════════════════════════════════════════════════════════════════════════

INFORMATIONS DU CANDIDAT À UTILISER :
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• Statut actuel : {academic_status}
• Entreprise visée : {company_name}
• Secteur/Domaine : {industry_sector if industry_sector else "Non précisé"}
• Localisation : {location if location else "Non précisé"}
• Stage/Poste recherché : {internship_type}
• Disponibilité : {availability}
• Expériences passées : {previous_experience}
• Contact/Networking : {networking_info if networking_info else "Non renseigné"}
• Qualités principales : {top_qualities}
• Fait marquant (actualité entreprise) : {striking_fact if striking_fact else "Non renseigné"}
• Projet futur/Side project : {future_plan if future_plan else "Non renseigné"}
• Rôle chez Révélation Entreprendre : {re_role}
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

STRUCTURE OBLIGATOIRE (4 parties) :

📌 PARAGRAPHE 1 - ACCROCHE & FIT CULTUREL
   - Commence par "Madame, Monsieur,"
   - Phrase d'intro : statut actuel + recherche de {internship_type} chez {company_name}
   - Ce qui t'attire chez {company_name} (positionnement, réputation, valeurs){sector_instruction}{networking_instruction}
   - Lien entre tes qualités ({top_qualities}) et la culture de l'entreprise{fact_instruction}

📌 PARAGRAPHE 2 - COMPÉTENCE TECHNIQUE
   - Décris ton expérience majeure : {previous_experience}
   - Ce que tu y as appris (rigueur, technique, méthodologie)
   - Montre ton implication (travail acharné, curiosité, initiative){future_instruction}

📌 PARAGRAPHE 3 - LEADERSHIP & EXCELLENCE
   - Mentionne ton excellence académique (classement, résultats)
   - Décris ton rôle chez Révélation Entreprendre : {re_role}
   - Rappelle les chiffres clés : plus grand concours de startups étudiant, 500+ startups candidates, 350 000€ de dotation
   - Fais le lien avec ce que tu apporteras à {company_name}

📌 CONCLUSION
   - Phrase de demande d'entretien (sobre et professionnelle)
   - Formule de politesse : "Dans l'attente du plaisir de vous rencontrer, je vous prie d'agréer, Madame, Monsieur, l'expression de mes respectueuses salutations."

RÈGLES IMPÉRATIVES :
━━━━━━━━━━━━━━━━━━━
✗ PAS d'en-tête (nom, email, téléphone)
✗ PAS de titre/objet
✗ PAS de signature à la fin
✓ Commence directement par "Madame, Monsieur,"
✓ Termine par la formule de politesse
✓ Longueur : environ 350-450 mots
✓ Ton : professionnel, confiant mais humble, dynamique
✓ Style : phrases fluides, pas de liste à puces, transitions naturelles
✓ ADAPTE le vocabulaire au secteur visé ({industry_sector if industry_sector else internship_type})

Génère maintenant le corps de la lettre :"""

    else:
        # ANGLAIS
        networking_instruction = ""
        if networking_info:
            networking_instruction = f"""
   - MUST INCLUDE this networking information: "{networking_info}"
   - Use it to justify your interest through insider feedback"""
        else:
            networking_instruction = """
   - No contact mentioned: reference the company's reputation or public sources"""

        fact_instruction = ""
        if striking_fact:
            fact_instruction = f"""
   - MUST MENTION this key fact: "{striking_fact}"
   - Integrate it naturally to show your knowledge of the company"""

        future_instruction = ""
        if future_plan:
            future_instruction = f"""
   - MENTION your future project: "{future_plan}" to show curiosity and ambition"""

        sector_instruction = ""
        if industry_sector:
            sector_instruction = f"""
   - Adapt vocabulary and references to the sector: {industry_sector}"""

        prompt = f"""You are an expert cover letter writer for professional applications.

OBJECTIVE: Write the BODY of a cover letter IN ENGLISH (from "Dear [Company] Recruitment Team," to the closing).

═══════════════════════════════════════════════════════════════════════════════
REFERENCE MODEL (imitate this style precisely)
═══════════════════════════════════════════════════════════════════════════════
{template_letter}
═══════════════════════════════════════════════════════════════════════════════

CANDIDATE INFORMATION TO USE:
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
• Current status: {academic_status}
• Target company: {company_name}
• Industry/Sector: {industry_sector if industry_sector else "Not specified"}
• Location: {location if location else "Not specified"}
• Position sought: {internship_type}
• Availability: {availability}
• Past experience: {previous_experience}
• Contact/Networking: {networking_info if networking_info else "Not provided"}
• Key qualities: {top_qualities}
• Notable fact (company news): {striking_fact if striking_fact else "Not provided"}
• Future project/Side project: {future_plan if future_plan else "Not provided"}
• Role at Révélation Entreprendre: {re_role}
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

MANDATORY STRUCTURE (4 parts):

📌 PARAGRAPH 1 - HOOK & CULTURAL FIT
   - Start with "Dear {company_name} Recruitment Team,"
   - Introduction: current status + seeking {internship_type} at {company_name}
   - What attracts you to {company_name} (positioning, reputation, values){sector_instruction}{networking_instruction}
   - Link between your qualities ({top_qualities}) and company culture{fact_instruction}

📌 PARAGRAPH 2 - TECHNICAL COMPETENCE
   - Describe your main experience: {previous_experience}
   - What you learned (rigor, technique, methodology)
   - Show your dedication (hard work, curiosity, initiative){future_instruction}

📌 PARAGRAPH 3 - LEADERSHIP & EXCELLENCE
   - Mention your academic excellence (ranking, results)
   - Describe your role at Révélation Entreprendre: {re_role}
   - Include key figures: largest student startup competition, 500+ startups, €350,000 in prizes
   - Connect to what you'll bring to {company_name}

📌 CONCLUSION
   - Request for interview (professional and concise)
   - Closing: "Yours sincerely,"

MANDATORY RULES:
━━━━━━━━━━━━━━━━━━━
✗ NO header (name, email, phone)
✗ NO subject line
✗ NO signature at the end
✓ Start directly with "Dear {company_name} Recruitment Team,"
✓ End with "Yours sincerely,"
✓ Length: approximately 300-400 words
✓ Tone: professional, confident yet humble, dynamic
✓ Style: fluid sentences, no bullet points, natural transitions
✓ ADAPT vocabulary to the target sector ({industry_sector if industry_sector else internship_type})

Generate the cover letter body now:"""

    return prompt


# ============================================================================
# TEMPLATE HTML - INTERFACE MODERNE MULTILINGUE
# ============================================================================
HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Générateur de Lettres | Révélation Entreprendre</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:wght@400;500;600;700&family=Inter:wght@300;400;500;600&display=swap" rel="stylesheet">
    <style>
        :root {
            --color-bg: #0a0a0a;
            --color-surface: #141414;
            --color-surface-elevated: #1a1a1a;
            --color-border: #2a2a2a;
            --color-border-hover: #3a3a3a;
            --color-text: #fafafa;
            --color-text-secondary: #888;
            --color-text-muted: #555;
            --color-accent: #c9a962;
            --color-accent-hover: #d4b87a;
            --color-accent-subtle: rgba(201, 169, 98, 0.1);
            --color-success: #22c55e;
            --color-error: #ef4444;
            --color-blue: #3b82f6;
            --font-display: 'Cormorant Garamond', Georgia, serif;
            --font-body: 'Inter', -apple-system, sans-serif;
            --radius-sm: 4px;
            --radius-md: 8px;
            --radius-lg: 12px;
            --shadow-sm: 0 1px 2px rgba(0,0,0,0.3);
            --shadow-md: 0 4px 12px rgba(0,0,0,0.4);
            --shadow-lg: 0 8px 30px rgba(0,0,0,0.5);
            --transition: 0.2s ease;
        }

        * {
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }

        html {
            scroll-behavior: smooth;
        }

        body {
            font-family: var(--font-body);
            background: var(--color-bg);
            color: var(--color-text);
            line-height: 1.6;
            min-height: 100vh;
        }

        /* Header */
        .header {
            position: fixed;
            top: 0;
            left: 0;
            right: 0;
            z-index: 100;
            padding: 1rem 2rem;
            background: rgba(10, 10, 10, 0.9);
            backdrop-filter: blur(20px);
            border-bottom: 1px solid var(--color-border);
        }

        .header-content {
            max-width: 1400px;
            margin: 0 auto;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }

        .logo {
            font-family: var(--font-display);
            font-size: 1.5rem;
            font-weight: 600;
            color: var(--color-text);
            text-decoration: none;
            letter-spacing: -0.02em;
        }

        .logo span {
            color: var(--color-accent);
        }

        .header-badge {
            font-size: 0.75rem;
            color: var(--color-text-secondary);
            padding: 0.25rem 0.75rem;
            border: 1px solid var(--color-border);
            border-radius: 20px;
        }

        /* Main Layout */
        .main {
            padding-top: 80px;
            min-height: 100vh;
        }

        .container {
            max-width: 1400px;
            margin: 0 auto;
            padding: 2rem;
        }

        /* Hero Section */
        .hero {
            text-align: center;
            padding: 4rem 0 3rem;
            position: relative;
        }

        .hero::before {
            content: '';
            position: absolute;
            top: 0;
            left: 50%;
            transform: translateX(-50%);
            width: 600px;
            height: 400px;
            background: radial-gradient(ellipse, rgba(201, 169, 98, 0.08) 0%, transparent 70%);
            pointer-events: none;
        }

        .hero-title {
            font-family: var(--font-display);
            font-size: clamp(2.5rem, 5vw, 4rem);
            font-weight: 500;
            letter-spacing: -0.03em;
            margin-bottom: 1rem;
            line-height: 1.1;
        }

        .hero-title span {
            color: var(--color-accent);
        }

        .hero-subtitle {
            font-size: 1.1rem;
            color: var(--color-text-secondary);
            max-width: 600px;
            margin: 0 auto;
        }

        /* Form Layout */
        .form-container {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 3rem;
            margin-top: 2rem;
        }

        @media (max-width: 1024px) {
            .form-container {
                grid-template-columns: 1fr;
            }
        }

        /* Form Panel */
        .form-panel {
            background: var(--color-surface);
            border: 1px solid var(--color-border);
            border-radius: var(--radius-lg);
            padding: 2rem;
        }

        .form-panel-header {
            margin-bottom: 2rem;
            padding-bottom: 1.5rem;
            border-bottom: 1px solid var(--color-border);
        }

        .form-panel-title {
            font-family: var(--font-display);
            font-size: 1.5rem;
            font-weight: 500;
            margin-bottom: 0.5rem;
        }

        .form-panel-subtitle {
            font-size: 0.875rem;
            color: var(--color-text-secondary);
        }

        /* Form Sections */
        .form-section {
            margin-bottom: 2rem;
        }

        .form-section:last-child {
            margin-bottom: 0;
        }

        .section-label {
            display: flex;
            align-items: center;
            gap: 0.5rem;
            font-size: 0.7rem;
            font-weight: 600;
            text-transform: uppercase;
            letter-spacing: 0.1em;
            color: var(--color-accent);
            margin-bottom: 1rem;
        }

        .section-label::after {
            content: '';
            flex: 1;
            height: 1px;
            background: var(--color-border);
        }

        /* Form Groups */
        .form-group {
            margin-bottom: 1.25rem;
        }

        .form-group:last-child {
            margin-bottom: 0;
        }

        .form-row {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 1rem;
        }

        .form-row-3 {
            display: grid;
            grid-template-columns: 1fr 1fr 1fr;
            gap: 1rem;
        }

        @media (max-width: 600px) {
            .form-row, .form-row-3 {
                grid-template-columns: 1fr;
            }
        }

        label {
            display: block;
            font-size: 0.875rem;
            font-weight: 500;
            color: var(--color-text);
            margin-bottom: 0.5rem;
        }

        label .required {
            color: var(--color-accent);
            margin-left: 2px;
        }

        label .optional {
            font-weight: 400;
            color: var(--color-text-muted);
            font-size: 0.8em;
        }

        /* Language Toggle */
        .language-toggle {
            display: flex;
            background: var(--color-bg);
            border: 1px solid var(--color-border);
            border-radius: var(--radius-md);
            padding: 4px;
            gap: 4px;
        }

        .language-toggle input[type="radio"] {
            display: none;
        }

        .language-toggle label {
            flex: 1;
            text-align: center;
            padding: 0.75rem 1.5rem;
            margin: 0;
            border-radius: var(--radius-sm);
            cursor: pointer;
            transition: var(--transition);
            font-size: 0.9rem;
            font-weight: 500;
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 0.5rem;
        }

        .language-toggle label:hover {
            background: var(--color-surface-elevated);
        }

        .language-toggle input[type="radio"]:checked + label {
            background: var(--color-accent);
            color: var(--color-bg);
        }

        .flag-icon {
            font-size: 1.2rem;
        }

        /* Inputs */
        input[type="text"],
        input[type="email"],
        input[type="password"],
        textarea,
        select {
            width: 100%;
            padding: 0.75rem 1rem;
            font-family: var(--font-body);
            font-size: 0.9375rem;
            color: var(--color-text);
            background: var(--color-bg);
            border: 1px solid var(--color-border);
            border-radius: var(--radius-md);
            transition: var(--transition);
            outline: none;
        }

        select {
            cursor: pointer;
            appearance: none;
            background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='12' height='12' viewBox='0 0 12 12'%3E%3Cpath fill='%23888' d='M6 8L1 3h10z'/%3E%3C/svg%3E");
            background-repeat: no-repeat;
            background-position: right 1rem center;
            padding-right: 2.5rem;
        }

        input:focus,
        textarea:focus,
        select:focus {
            border-color: var(--color-accent);
            box-shadow: 0 0 0 3px var(--color-accent-subtle);
        }

        input::placeholder,
        textarea::placeholder {
            color: var(--color-text-muted);
        }

        textarea {
            resize: vertical;
            min-height: 100px;
        }

        /* Helper text */
        .helper-text {
            font-size: 0.8rem;
            color: var(--color-text-muted);
            margin-top: 0.4rem;
        }

        /* Buttons */
        .btn {
            display: inline-flex;
            align-items: center;
            justify-content: center;
            gap: 0.5rem;
            padding: 0.875rem 2rem;
            font-family: var(--font-body);
            font-size: 0.9375rem;
            font-weight: 500;
            border: none;
            border-radius: var(--radius-md);
            cursor: pointer;
            transition: var(--transition);
            text-decoration: none;
        }

        .btn-primary {
            width: 100%;
            background: var(--color-accent);
            color: var(--color-bg);
        }

        .btn-primary:hover {
            background: var(--color-accent-hover);
            transform: translateY(-1px);
            box-shadow: var(--shadow-md);
        }

        .btn-primary:active {
            transform: translateY(0);
        }

        .btn-primary:disabled {
            opacity: 0.6;
            cursor: not-allowed;
            transform: none;
        }

        .btn-secondary {
            background: transparent;
            color: var(--color-accent);
            border: 1px solid var(--color-accent);
        }

        .btn-secondary:hover {
            background: var(--color-accent-subtle);
        }

        .btn-download {
            background: var(--color-success);
            color: white;
        }

        .btn-download:hover {
            filter: brightness(1.1);
        }

        .btn-pdf {
            background: var(--color-error);
            color: white;
        }

        .btn-pdf:hover {
            filter: brightness(1.1);
        }

        /* Output Panel */
        .output-panel {
            background: var(--color-surface);
            border: 1px solid var(--color-border);
            border-radius: var(--radius-lg);
            overflow: hidden;
            display: flex;
            flex-direction: column;
        }

        .output-header {
            padding: 1.5rem 2rem;
            border-bottom: 1px solid var(--color-border);
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 1rem;
        }

        .output-title {
            font-family: var(--font-display);
            font-size: 1.25rem;
            font-weight: 500;
        }

        .output-actions {
            display: flex;
            gap: 0.75rem;
        }

        .output-content {
            padding: 2rem;
            flex: 1;
            overflow-y: auto;
            max-height: calc(100vh - 300px);
        }

        .output-placeholder {
            display: flex;
            flex-direction: column;
            align-items: center;
            justify-content: center;
            text-align: center;
            padding: 4rem 2rem;
            color: var(--color-text-muted);
        }

        .output-placeholder-icon {
            font-size: 3rem;
            margin-bottom: 1rem;
            opacity: 0.5;
        }

        .output-placeholder-text {
            font-size: 1rem;
            margin-bottom: 0.5rem;
        }

        .output-placeholder-hint {
            font-size: 0.875rem;
        }

        .letter-preview {
            font-family: var(--font-body);
            font-size: 0.9375rem;
            line-height: 1.8;
            white-space: pre-wrap;
            color: var(--color-text);
        }

        /* Language badge */
        .lang-badge {
            display: inline-flex;
            align-items: center;
            gap: 0.3rem;
            font-size: 0.75rem;
            padding: 0.25rem 0.5rem;
            background: var(--color-surface-elevated);
            border: 1px solid var(--color-border);
            border-radius: 4px;
            margin-left: 0.5rem;
        }

        /* Alerts */
        .alert {
            padding: 1rem 1.25rem;
            border-radius: var(--radius-md);
            margin-bottom: 1.5rem;
            font-size: 0.9rem;
        }

        .alert-error {
            background: rgba(239, 68, 68, 0.1);
            border: 1px solid rgba(239, 68, 68, 0.3);
            color: #fca5a5;
        }

        .alert-success {
            background: rgba(34, 197, 94, 0.1);
            border: 1px solid rgba(34, 197, 94, 0.3);
            color: #86efac;
        }

        /* Loading State */
        .spinner {
            width: 20px;
            height: 20px;
            border: 2px solid var(--color-border);
            border-top-color: var(--color-bg);
            border-radius: 50%;
            animation: spin 0.8s linear infinite;
            display: inline-block;
        }

        @keyframes spin {
            to { transform: rotate(360deg); }
        }

        /* Copy button */
        .btn-copy {
            padding: 0.5rem 1rem;
            font-size: 0.8rem;
            background: var(--color-surface-elevated);
            border: 1px solid var(--color-border);
            color: var(--color-text);
        }

        .btn-copy:hover {
            border-color: var(--color-accent);
        }

        .btn-copy.copied {
            background: var(--color-success);
            border-color: var(--color-success);
            color: white;
        }

        /* Footer */
        .footer {
            text-align: center;
            padding: 2rem;
            color: var(--color-text-muted);
            font-size: 0.8rem;
            border-top: 1px solid var(--color-border);
            margin-top: 4rem;
        }

        /* Animations */
        @keyframes fadeIn {
            from { opacity: 0; transform: translateY(10px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .fade-in {
            animation: fadeIn 0.4s ease forwards;
        }

        /* Sticky output on desktop */
        @media (min-width: 1025px) {
            .output-panel {
                position: sticky;
                top: 100px;
                max-height: calc(100vh - 120px);
            }
        }

        /* Sector chips */
        .sector-suggestions {
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem;
            margin-top: 0.5rem;
        }

        .sector-chip {
            font-size: 0.75rem;
            padding: 0.25rem 0.75rem;
            background: var(--color-surface-elevated);
            border: 1px solid var(--color-border);
            border-radius: 20px;
            cursor: pointer;
            transition: var(--transition);
        }

        .sector-chip:hover {
            border-color: var(--color-accent);
            color: var(--color-accent);
        }
    </style>
</head>
<body>
    <header class="header">
        <div class="header-content">
            <a href="/" class="logo">Révélation <span>Entreprendre</span></a>
            <span class="header-badge">Cover Letter Generator v4</span>
        </div>
    </header>

    <main class="main">
        <div class="container">
            <section class="hero">
                <h1 class="hero-title">Générateur de <span>Lettres</span></h1>
                <p class="hero-subtitle">Créez une lettre de motivation personnalisée en français ou en anglais, adaptée à tous les secteurs.</p>
            </section>

            {% if error %}
                <div class="alert alert-error fade-in">
                    ⚠️ {{ error }}
                </div>
            {% endif %}

            {% if success_message %}
                <div class="alert alert-success fade-in">
                    ✓ {{ success_message }}
                </div>
            {% endif %}

            <form method="POST" id="letterForm">
                <div class="form-container">
                    <!-- Left Panel: Form -->
                    <div class="form-panel">
                        <div class="form-panel-header">
                            <h2 class="form-panel-title">Informations</h2>
                            <p class="form-panel-subtitle">Remplissez les champs pour générer votre lettre personnalisée</p>
                        </div>

                        <!-- Language Selection -->
                        <div class="form-section">
                            <div class="section-label">Langue / Language</div>
                            <div class="language-toggle">
                                <input type="radio" name="language" id="lang-fr" value="fr" {{ 'checked' if form.get('language', 'fr') == 'fr' else '' }}>
                                <label for="lang-fr">
                                    <span class="flag-icon">🇫🇷</span>
                                    Français
                                </label>
                                <input type="radio" name="language" id="lang-en" value="en" {{ 'checked' if form.get('language') == 'en' else '' }}>
                                <label for="lang-en">
                                    <span class="flag-icon">🇬🇧</span>
                                    English
                                </label>
                            </div>
                        </div>

                        <!-- Access Section -->
                        <div class="form-section">
                            <div class="section-label">Accès</div>
                            <div class="form-group">
                                <label>Code d'accès RE <span class="required">*</span></label>
                                <input type="password" name="secret_code" value="{{ form.get('secret_code', '') }}" required placeholder="Entrez le code Révélation Entreprendre">
                            </div>
                        </div>

                        <!-- Personal Info Section -->
                        <div class="form-section">
                            <div class="section-label">Coordonnées</div>
                            <div class="form-row">
                                <div class="form-group">
                                    <label>Nom complet <span class="required">*</span></label>
                                    <input type="text" name="candidate_name" value="{{ form.get('candidate_name', '') }}" required placeholder="Prénom Nom">
                                </div>
                                <div class="form-group">
                                    <label>Téléphone <span class="required">*</span></label>
                                    <input type="text" name="candidate_phone" value="{{ form.get('candidate_phone', '') }}" required placeholder="+33 6 XX XX XX XX">
                                </div>
                            </div>
                            <div class="form-group">
                                <label>Email <span class="required">*</span></label>
                                <input type="email" name="candidate_email" value="{{ form.get('candidate_email', '') }}" required placeholder="prenom.nom@edhec.com">
                            </div>
                            <div class="form-row">
                                <div class="form-group">
                                    <label>Lieu et date</label>
                                    <input type="text" name="city_and_date" value="{{ form.get('city_and_date', today_date) }}" placeholder="Fait à Lille, le ...">
                                </div>
                                <div class="form-group">
                                    <label>Statut académique</label>
                                    <input type="text" name="academic_status" value="{{ form.get('academic_status', '') }}" placeholder="M1 Finance Track EDHEC">
                                </div>
                            </div>
                        </div>

                        <!-- Target Section -->
                        <div class="form-section">
                            <div class="section-label">Candidature</div>
                            <div class="form-row">
                                <div class="form-group">
                                    <label>Entreprise visée <span class="required">*</span></label>
                                    <input type="text" name="company_name" value="{{ form.get('company_name', '') }}" required placeholder="Ex: McKinsey, L'Oréal, BNP...">
                                </div>
                                <div class="form-group">
                                    <label>Localisation <span class="optional">(optionnel)</span></label>
                                    <input type="text" name="location" value="{{ form.get('location', '') }}" placeholder="Ex: Paris, London, NYC...">
                                </div>
                            </div>
                            <div class="form-group">
                                <label>Secteur / Domaine <span class="optional">(optionnel)</span></label>
                                <input type="text" name="industry_sector" id="industry_sector" value="{{ form.get('industry_sector', '') }}" placeholder="Ex: M&A, Conseil en stratégie, Marketing, Tech...">
                                <div class="sector-suggestions">
                                    <span class="sector-chip" onclick="setSector('M&A / Investment Banking')">M&A</span>
                                    <span class="sector-chip" onclick="setSector('Private Equity')">Private Equity</span>
                                    <span class="sector-chip" onclick="setSector('Conseil en stratégie')">Conseil</span>
                                    <span class="sector-chip" onclick="setSector('Audit / Transaction Services')">Audit/TS</span>
                                    <span class="sector-chip" onclick="setSector('Marketing / Brand Management')">Marketing</span>
                                    <span class="sector-chip" onclick="setSector('Tech / Product Management')">Tech/Product</span>
                                    <span class="sector-chip" onclick="setSector('Asset Management')">Asset Mgmt</span>
                                </div>
                            </div>
                            <div class="form-row">
                                <div class="form-group">
                                    <label>Type de stage/poste <span class="required">*</span></label>
                                    <input type="text" name="internship_type" value="{{ form.get('internship_type', '') }}" required placeholder="Ex: stage en M&A, CDI consultant...">
                                </div>
                                <div class="form-group">
                                    <label>Disponibilité <span class="required">*</span></label>
                                    <input type="text" name="availability" value="{{ form.get('availability', '') }}" required placeholder="Ex: September 2026">
                                </div>
                            </div>
                        </div>

                        <!-- Experience Section -->
                        <div class="form-section">
                            <div class="section-label">Expérience & Compétences</div>
                            <div class="form-group">
                                <label>Expériences professionnelles <span class="required">*</span></label>
                                <textarea name="previous_experience" rows="4" required placeholder="Décrivez vos expériences : poste, entreprise, missions, apprentissages...">{{ form.get('previous_experience', '') }}</textarea>
                            </div>
                            <div class="form-group">
                                <label>Top 3 qualités <span class="required">*</span></label>
                                <textarea name="top_qualities" rows="3" required placeholder="Ex: Rigueur (prépa), leadership (capitaine équipe), analyse (projets)...">{{ form.get('top_qualities', '') }}</textarea>
                            </div>
                            <div class="form-group">
                                <label>Rôle chez Révélation Entreprendre <span class="required">*</span></label>
                                <input type="text" name="re_role" value="{{ form.get('re_role', '') }}" required placeholder="Ex: Responsable Partenariats, VP Events...">
                            </div>
                        </div>

                        <!-- Optional Section -->
                        <div class="form-section">
                            <div class="section-label">Personnalisation (optionnel)</div>
                            <div class="form-group">
                                <label>Contact / Networking <span class="optional">(recommandé)</span></label>
                                <textarea name="networking_info" rows="2" placeholder="Ex: Discussion with John Smith, Analyst, during EDHEC Finance Forum...">{{ form.get('networking_info', '') }}</textarea>
                            </div>
                            <div class="form-group">
                                <label>Fait marquant sur l'entreprise <span class="optional">(recommandé)</span></label>
                                <textarea name="striking_fact" rows="2" placeholder="Ex: Recent deal X, company culture, specific project...">{{ form.get('striking_fact', '') }}</textarea>
                            </div>
                            <div class="form-group">
                                <label>Projet futur / Side project</label>
                                <textarea name="future_plan" rows="2" placeholder="Ex: Exchange semester at LSE, personal project, career goal...">{{ form.get('future_plan', '') }}</textarea>
                            </div>
                        </div>

                        <button type="submit" class="btn btn-primary" id="submitBtn">
                            <svg width="20" height="20" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">
                                <path d="M12 20h9M16.5 3.5a2.121 2.121 0 0 1 3 3L7 19l-4 1 1-4L16.5 3.5z"/>
                            </svg>
                            Générer la lettre
                        </button>
                    </div>

                    <!-- Right Panel: Output -->
                    <div class="output-panel">
                        <div class="output-header">
                            <h2 class="output-title">
                                Lettre générée
                                {% if letter %}
                                    <span class="lang-badge">
                                        {% if session.get('language') == 'en' %}🇬🇧 EN{% else %}🇫🇷 FR{% endif %}
                                    </span>
                                {% endif %}
                            </h2>
                            {% if letter %}
                            <div class="output-actions">
                                <button type="button" class="btn btn-copy" onclick="copyToClipboard()">
                                    <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">
                                        <rect x="9" y="9" width="13" height="13" rx="2" ry="2"></rect>
                                        <path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"></path>
                                    </svg>
                                    Copier
                                </button>
                                <a href="/download/docx" class="btn btn-download">
                                    <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">
                                        <path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4M7 10l5 5 5-5M12 15V3"/>
                                    </svg>
                                    Word
                                </a>
                            </div>
                            {% endif %}
                        </div>
                        <div class="output-content">
                            {% if letter %}
                                <div class="letter-preview fade-in" id="letterContent">{{ letter }}</div>
                            {% else %}
                                <div class="output-placeholder">
                                    <div class="output-placeholder-icon">✉️</div>
                                    <p class="output-placeholder-text">Votre lettre apparaîtra ici</p>
                                    <p class="output-placeholder-hint">Remplissez le formulaire et cliquez sur "Générer"</p>
                                </div>
                            {% endif %}
                        </div>
                    </div>
                </div>
            </form>
        </div>
    </main>

    <footer class="footer">
        <p>© {{ current_year }} Révélation Entreprendre – EDHEC Business School</p>
        <p style="margin-top: 0.5rem;">Générateur propulsé par l'IA • FR / EN • Tous secteurs</p>
    </footer>

    <script>
        // Set sector from chip click
        function setSector(sector) {
            document.getElementById('industry_sector').value = sector;
        }

        // Copy to clipboard
        function copyToClipboard() {
            const letterContent = document.getElementById('letterContent');
            if (!letterContent) return;
            
            navigator.clipboard.writeText(letterContent.innerText).then(() => {
                const btn = document.querySelector('.btn-copy');
                const originalHTML = btn.innerHTML;
                btn.innerHTML = '✓ Copié !';
                btn.classList.add('copied');
                setTimeout(() => {
                    btn.innerHTML = originalHTML;
                    btn.classList.remove('copied');
                }, 2000);
            });
        }

        // Form submission loading state
        document.getElementById('letterForm').addEventListener('submit', function() {
            const btn = document.getElementById('submitBtn');
            btn.disabled = true;
            btn.innerHTML = '<span class="spinner"></span> Génération en cours...';
        });

        // Update placeholders based on language
        document.querySelectorAll('input[name="language"]').forEach(radio => {
            radio.addEventListener('change', function() {
                const isEn = this.value === 'en';
                // Update some placeholders
                document.querySelector('[name="city_and_date"]').placeholder = isEn ? 'London, January 4th, 2026' : 'Fait à Lille, le 4 janvier 2026';
                document.querySelector('[name="academic_status"]').placeholder = isEn ? 'Master in Finance at EDHEC' : 'M1 Finance Track EDHEC';
                document.querySelector('[name="internship_type"]').placeholder = isEn ? 'M&A Analyst Internship' : 'stage en M&A';
                document.querySelector('[name="availability"]').placeholder = isEn ? 'September 2026' : 'septembre 2026';
            });
        });

        // Smooth scroll to output on mobile after generation
        {% if letter %}
        if (window.innerWidth < 1024) {
            document.querySelector('.output-panel').scrollIntoView({ behavior: 'smooth', block: 'start' });
        }
        {% endif %}
    </script>
</body>
</html>
"""


# ============================================================================
# FONCTIONS UTILITAIRES
# ============================================================================
def replace_placeholder(doc: Document, placeholder: str, new_text: str):
    """Remplace un placeholder dans le document Word."""
    for p in doc.paragraphs:
        if placeholder in p.text:
            new_para_text = p.text.replace(placeholder, new_text)
            if p.runs:
                p.runs[0].text = new_para_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new_para_text)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if placeholder in p.text:
                        new_para_text = p.text.replace(placeholder, new_text)
                        if p.runs:
                            p.runs[0].text = new_para_text
                            for r in p.runs[1:]:
                                r.text = ""
                        else:
                            p.add_run(new_para_text)


def replace_body_with_indentation(doc: Document, placeholder: str, body_text: str, indent_cm: float = 1.0, space_after_pt: float = 12.0):
    """
    Remplace le placeholder du corps par le texte avec alinéas au début de chaque paragraphe.
    
    Args:
        doc: Document Word
        placeholder: Le placeholder à remplacer (ex: {{corps}})
        body_text: Le texte du corps de la lettre
        indent_cm: Taille de l'alinéa en centimètres (défaut: 1.0 cm)
        space_after_pt: Espace après chaque paragraphe en points (défaut: 12.0 pt)
    """
    from docx.shared import Cm, Pt
    
    # Sépare le corps en paragraphes
    paragraphs = [p.strip() for p in body_text.split('\n\n') if p.strip()]
    
    for p in doc.paragraphs:
        if placeholder in p.text:
            # Récupère le style et le format du paragraphe original
            original_style = p.style
            
            # Remplace le premier paragraphe dans le placeholder existant
            if paragraphs:
                first_para = paragraphs[0]
                p.clear()
                p.add_run(first_para)
                p.paragraph_format.first_line_indent = Cm(indent_cm)
                p.paragraph_format.space_after = Pt(space_after_pt)
                
                # Ajoute les paragraphes suivants après celui-ci
                parent = p._element.getparent()
                index = list(parent).index(p._element)
                
                for i, para_text in enumerate(paragraphs[1:], start=1):
                    # Crée un nouveau paragraphe
                    new_p = doc.add_paragraph(para_text)
                    new_p.style = original_style
                    new_p.paragraph_format.first_line_indent = Cm(indent_cm)
                    new_p.paragraph_format.space_after = Pt(space_after_pt)
                    
                    # Déplace le paragraphe à la bonne position
                    parent.insert(index + i, new_p._element)
            
            return  # On a trouvé et remplacé, on sort


def get_today_date(language="fr"):
    """Retourne la date du jour formatée selon la langue."""
    today = datetime.now()
    
    if language == "en":
        # Format anglais: London, January 4th, 2026
        day = today.day
        suffix = "th" if 11 <= day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
        months_en = ["January", "February", "March", "April", "May", "June",
                     "July", "August", "September", "October", "November", "December"]
        return f"London, {months_en[today.month-1]} {day}{suffix}, {today.year}"
    else:
        # Format français: Fait à Lille, le 4 janvier 2026
        months_fr = {
            1: "janvier", 2: "février", 3: "mars", 4: "avril",
            5: "mai", 6: "juin", 7: "juillet", 8: "août",
            9: "septembre", 10: "octobre", 11: "novembre", 12: "décembre"
        }
        return f"Fait à Lille, le {today.day} {months_fr[today.month]} {today.year}"


# ============================================================================
# ROUTES
# ============================================================================
@app.route("/", methods=["GET", "POST"])
def index():
    letter = None
    error = None
    success_message = None
    
    if request.method == "POST":
        secret_code = request.form.get("secret_code", "").strip()
        language = request.form.get("language", "fr").strip()
        
        if secret_code != RE_SECRET_CODE:
            error = "Code d'accès incorrect. Cette application est réservée aux membres de Révélation Entreprendre."
        else:
            # Sauvegarde des données en session
            session["language"] = language
            session["candidate_name"] = request.form.get("candidate_name", "").strip()
            session["candidate_email"] = request.form.get("candidate_email", "").strip()
            session["candidate_phone"] = request.form.get("candidate_phone", "").strip()
            session["city_and_date"] = request.form.get("city_and_date", "").strip()
            session["company_name"] = request.form.get("company_name", "").strip()
            session["internship_type"] = request.form.get("internship_type", "").strip()
            session["availability"] = request.form.get("availability", "").strip()
            session["location"] = request.form.get("location", "").strip()
            session["industry_sector"] = request.form.get("industry_sector", "").strip()
            
            # Construction et envoi du prompt
            prompt = build_prompt(request.form)
            
            try:
                response = client.chat.completions.create(
                    model="gpt-4o",
                    messages=[
                        {
                            "role": "system",
                            "content": f"""You are an expert cover letter writer. 
Write {'in English' if language == 'en' else 'in French'}.
Create professional, fluent, and convincing letters.
Strictly follow the requested structure and naturally integrate the provided information.
Adapt the vocabulary and tone to the target industry/sector."""
                        },
                        {
                            "role": "user",
                            "content": prompt
                        }
                    ],
                    temperature=0.7,
                    max_tokens=2000
                )
                
                letter = response.choices[0].message.content.strip()
                
                # Nettoyage
                letter = re.sub(r'^["\']|["\']$', '', letter)
                letter = re.sub(r'^```\w*\n?|```$', '', letter, flags=re.MULTILINE)
                
                session["letter_body"] = letter
                success_message = "Lettre générée avec succès !" if language == "fr" else "Cover letter generated successfully!"
                
            except Exception as e:
                error = f"Erreur lors de la génération : {str(e)}"
    
    # Date par défaut selon la langue sélectionnée
    selected_lang = request.form.get("language", "fr")
    
    return render_template_string(
        HTML_TEMPLATE,
        letter=letter,
        error=error,
        success_message=success_message,
        form=request.form,
        today_date=get_today_date(selected_lang),
        current_year=datetime.now().year,
        session=session
    )


@app.route("/download/docx")
def download_docx():
    """Génère et télécharge la lettre au format Word en utilisant le template avec placeholders."""
    letter_body = session.get("letter_body")
    
    if not letter_body:
        return "Aucune lettre à télécharger. Veuillez d'abord générer une lettre.", 400
    
    language = session.get("language", "fr")
    candidate_name = session.get("candidate_name", "")
    candidate_email = session.get("candidate_email", "")
    candidate_phone = session.get("candidate_phone", "")
    city_and_date = session.get("city_and_date", "")
    company_name = session.get("company_name", "")
    internship_type = session.get("internship_type", "")
    availability = session.get("availability", "")
    
    # Construction du titre selon la langue
    if language == "en":
        if availability:
            titre = f"Application for {internship_type} at {company_name} starting in {availability}"
        else:
            titre = f"Application for {internship_type} at {company_name}"
    else:
        if availability:
            titre = f"Candidature pour un {internship_type} chez {company_name} à partir de {availability}"
        else:
            titre = f"Candidature pour un {internship_type} chez {company_name}"
    
    # Vérification du template
    if not os.path.exists(TEMPLATE_DOCX_PATH):
        return "Template Word introuvable. Placez 'template placeholder.docx' dans le dossier.", 400
    
    # Chargement et remplacement des placeholders
    doc = Document(TEMPLATE_DOCX_PATH)
    
    replace_placeholder(doc, "{{nom}}", candidate_name)
    replace_placeholder(doc, "{{mail}}", candidate_email)
    replace_placeholder(doc, "{{tel}}", candidate_phone)
    replace_placeholder(doc, "{{entreprise}}", company_name)
    replace_placeholder(doc, "{{faità}}", city_and_date)
    replace_placeholder(doc, "{{titre}}", titre)
    
    # Corps avec alinéas (1 cm par défaut, ajuste si besoin)
    replace_body_with_indentation(doc, "{{corps}}", letter_body, indent_cm=1.0)
    
    replace_placeholder(doc, "{{signature}}", candidate_name)
    
    # Sauvegarde en mémoire
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Nom du fichier
    safe_company = re.sub(r'[^\w\s-]', '', company_name).strip().replace(' ', '_')
    lang_suffix = "EN" if language == "en" else "FR"
    filename = f"CoverLetter_{safe_company}_{lang_suffix}_{datetime.now().strftime('%Y%m%d')}.docx"
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


# ============================================================================
# POINT D'ENTRÉE
# ============================================================================
if __name__ == "__main__":
    if not os.getenv("OPENAI_API_KEY"):
        print("⚠️  ATTENTION: Variable OPENAI_API_KEY non définie !")
        print("   Créez un fichier .env avec : OPENAI_API_KEY=sk-votre-cle-ici")
    
    print("🚀 Démarrage du serveur...")
    print("   URL: http://localhost:8080")
    print("   Langues: FR 🇫🇷 / EN 🇬🇧")
    print("   Secteurs: Tous domaines supportés")
    app.run(debug=True, host="0.0.0.0", port=8080)
